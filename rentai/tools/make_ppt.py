#!/usr/bin/env python
# -*- coding: utf-8 -*-

"""
임대차계약서 이상탐지 시스템 발표용 PPT 생성기

사용 예:
  python rentai/tools/make_ppt.py \
    --project "임대차계약서 이상탐지 시스템" \
    --presenter "박병현" \
    --out "docs/임대차계약서_이상탐지_시스템_발표.pptx"
"""

from __future__ import annotations

import argparse
from datetime import datetime
from pathlib import Path
from typing import List

from pptx import Presentation
from pptx.util import Inches, Pt


def add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    slide_layout = prs.slide_layouts[0]  # Title Slide
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def add_bullet_slide(
    prs: Presentation,
    title: str,
    bullets: List[str],
    notes: str = "",
) -> None:
    slide_layout = prs.slide_layouts[1]  # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title

    body = slide.placeholders[1].text_frame
    if not bullets:
        body.text = ""
    else:
        body.text = bullets[0]
        for b in bullets[1:]:
            p = body.add_paragraph()
            p.text = b
            p.level = 0

    if notes:
        notes_slide = slide.notes_slide
        notes_slide.notes_text_frame.text = notes


def build_deck(project: str, presenter: str) -> Presentation:
    today = datetime.now().strftime("%Y-%m-%d")
    prs = Presentation()

    # 1) Title
    add_title_slide(
        prs,
        title=project,
        subtitle=f"발표자: {presenter} | 날짜: {today}",
    )

    # 2) 프로젝트 개요
    add_bullet_slide(
        prs,
        title="프로젝트 개요",
        bullets=[
            "목적: 전세계약서 PDF 자동 분석·이상 탐지·보고서 생성",
            "구성: FastAPI 백엔드, Streamlit 프론트, RAG(Chroma), ReportLab",
            "이번 범위: 정리/통합/정합화 및 RAG 인덱스 재구축",
        ],
        notes=(
            "이번 스프린트는 코드 정리와 실행 경로 통일, 응답 스키마 정합화에 집중했습니다."
        ),
    )

    # 3) 현재 아키텍처 요약
    add_bullet_slide(
        prs,
        title="현재 아키텍처",
        bullets=[
            "백엔드 진입점 단일화: backend/fastapi_app.py",
            "서비스: 파서/필드추출/룰엔진/RAG/리포트",
            "프론트: frontend/app.py (API 연동 및 보고서 다운로드)",
        ],
        notes="fastapi_app.py가 허브 역할을 합니다.",
    )

    # 4) 정리 작업(클린업)
    add_bullet_slide(
        prs,
        title="정리 작업(클린업)",
        bullets=[
            "__pycache__/ .pyc, 미사용 템플릿 삭제",
            "레거시 파일은 '미사용/레거시' 주석 표시",
            "효과: 혼선 제거, 유지보수성 향상",
        ],
        notes="불필요 파일을 제거하고 레거시를 명시했습니다.",
    )

    # 5) 백엔드 진입점 통일
    add_bullet_slide(
        prs,
        title="백엔드 진입점 통일",
        bullets=[
            "app/main.py ↔ fastapi_app.py 중복 제거",
            "공식 엔드포인트는 fastapi_app.py에 집중",
            "장점: 단일 실행/배포 경로, 디버깅 단순화",
        ],
        notes="단일 진입점으로 운영 리스크를 낮췄습니다.",
    )

    # 6) API 응답 스키마 표준화
    add_bullet_slide(
        prs,
        title="API 응답 스키마 표준화",
        bullets=[
            "프론트 기대 구조: summary / fields / issues / rag",
            "하위 호환: 원본 데이터 동시 제공",
            "프론트 렌더링 로직 간소화",
        ],
        notes="/analyze/pdf 응답을 표준화했습니다.",
    )

    # 7) 레거시 라우터 정리
    add_bullet_slide(
        prs,
        title="레거시 라우터 정리",
        bullets=[
            "app/api/analyze.py → 410 Gone(안내 전용)",
            "app/main.py에서 레거시 include 제거",
            "공식 엔드포인트 안내 명확화",
        ],
        notes="공식 경로만 사용하도록 유도했습니다.",
    )

    # 8) RAG 인덱스 재구축
    add_bullet_slide(
        prs,
        title="RAG 인덱스 재구축",
        bullets=[
            "tools/ingest.py로 임베딩 생성, vectorstore/chroma 저장",
            "/kb/status, /reference/law 점검",
            "말뭉치: lease_law_samples.txt 등",
        ],
        notes="상태 확인과 간단 검색으로 점검 가능합니다.",
    )

    # 9) 데모 플로우
    add_bullet_slide(
        prs,
        title="데모 플로우",
        bullets=[
            "업로드 → /analyze/pdf → 결과 요약/이슈/RAG 근거",
            "보고서 다운로드: /report/pdf (원본 PDF 첨부)",
            "PDF: 필드/이슈/근거/프리뷰 포함",
        ],
        notes="엔드투엔드 데모가 가능합니다.",
    )

    # 10) 변경 전/후 요약
    add_bullet_slide(
        prs,
        title="변경 전/후 요약",
        bullets=[
            "전: 진입점 이원화, 응답 스키마 불일치, 레거시 혼재",
            "후: 단일 진입점, 표준 스키마, 레거시 명시/정리",
        ],
        notes="통일성과 일관성을 확보했습니다.",
    )

    # 11) 리스크 및 대응
    add_bullet_slide(
        prs,
        title="리스크 및 대응",
        bullets=[
            "RAG 품질: 말뭉치 확대·정제",
            "파서 정밀도: 규칙/정규식 개선",
            "키 관리: .env 및 권한 관리",
        ],
        notes="남은 리스크와 대응 방향을 제시합니다.",
    )

    # 12) 다음 계획(로드맵)
    add_bullet_slide(
        prs,
        title="다음 계획",
        bullets=[
            "테스트 보강(파서/룰엔진)",
            "예외 처리/로깅 개선",
            "PDF 레이아웃 개선, 도커/CI 도입",
        ],
        notes="다음 스프린트 액션 아이템입니다.",
    )

    # 13) Q&A
    add_bullet_slide(
        prs,
        title="Q&A",
        bullets=["질문을 받겠습니다."],
        notes="데모, 한계점, 확장 계획 등 Q&A",  # 간단 노트
    )

    return prs


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--project", required=True)
    ap.add_argument("--presenter", required=True)
    ap.add_argument("--out", default="docs/presentation_rentai.pptx")
    args = ap.parse_args()

    prs = build_deck(project=args.project, presenter=args.presenter)

    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    print(f"[PPT OK] saved -> {out_path}")


if __name__ == "__main__":
    main()


